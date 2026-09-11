"""
Packing List Excel Service.

Parses, validates, and generates Excel packing lists (.xlsx) for suppliers.
Adheres to the 12-column factory packing sheet format:
  Col A: P/O #
  Col B: PO item
  Col C: Pack No.
  Col D: Cart No
  Col E: Supplier_Carton_ref
  Col F: ProductCode
  Col G: Lot No.
  Col H: Width
  Col I: GW (Gross Weight)
  Col J: NW (Net Weight)
  Col K: Quantity
  Col L: UOM (Unit of Measure)
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.purchase_order import PurchaseOrder


@dataclass
class ParsedCartonRow:
    row_index: int
    po_number: str
    po_item: str
    pack_number: str
    carton_number: int
    supplier_carton_ref: str
    product_code: str
    lot_number: str
    width: Optional[float]
    gross_weight: float
    net_weight: float
    quantity: float
    uom: str
    errors: List[str] = field(default_factory=list)


@dataclass
class LineItemSummary:
    po_number: str
    po_item: str
    product_code: str
    total_qty: float
    carton_count: int
    total_gw: float
    total_nw: float
    po_ordered_qty: Optional[float] = None
    po_open_balance: Optional[float] = None
    is_valid: bool = True
    errors: List[str] = field(default_factory=list)


@dataclass
class ExcelValidationResult:
    is_valid: bool
    total_rows: int
    total_cartons: int
    total_gross_weight: float
    total_net_weight: float
    total_quantity: float
    line_summaries: List[LineItemSummary]
    rows: List[ParsedCartonRow]
    general_errors: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "is_valid": self.is_valid,
            "total_rows": self.total_rows,
            "total_cartons": self.total_cartons,
            "total_gross_weight": round(self.total_gross_weight, 2),
            "total_net_weight": round(self.total_net_weight, 2),
            "total_quantity": round(self.total_quantity, 2),
            "general_errors": self.general_errors,
            "line_summaries": [
                {
                    "po_number": s.po_number,
                    "po_item": s.po_item,
                    "product_code": s.product_code,
                    "total_qty": round(s.total_qty, 2),
                    "carton_count": s.carton_count,
                    "total_gw": round(s.total_gw, 2),
                    "total_nw": round(s.total_nw, 2),
                    "po_ordered_qty": s.po_ordered_qty,
                    "po_open_balance": s.po_open_balance,
                    "is_valid": s.is_valid,
                    "errors": s.errors,
                }
                for s in self.line_summaries
            ],
            "rows": [
                {
                    "row_index": r.row_index,
                    "po_number": r.po_number,
                    "po_item": r.po_item,
                    "pack_number": r.pack_number,
                    "carton_number": r.carton_number,
                    "supplier_carton_ref": r.supplier_carton_ref,
                    "product_code": r.product_code,
                    "lot_number": r.lot_number,
                    "width": r.width,
                    "gross_weight": r.gross_weight,
                    "net_weight": r.net_weight,
                    "quantity": r.quantity,
                    "uom": r.uom,
                    "errors": r.errors,
                }
                for r in self.rows
            ],
        }


def _clean_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _parse_float(val: Any) -> float:
    if val is None:
        return 0.0
    val_str = str(val).strip().replace(",", ".")
    try:
        return float(val_str)
    except ValueError:
        return 0.0


def _parse_int(val: Any) -> int:
    if val is None:
        return 0
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return 0


def _normalize_item_num(val: Any) -> str:
    raw = _clean_str(val)
    if not raw:
        return ""
    # Strip schedule suffix if present e.g. 00100-0001 -> 00100
    if "-" in raw:
        raw = raw.split("-")[0]
    # Pad to 5 digits if purely numeric e.g. 100 -> 00100
    if raw.isdigit():
        return raw.zfill(5)
    return raw


async def parse_and_validate_packing_excel(
    file_bytes: bytes,
    supplier_id: Optional[str] = None,
    supplier_code: Optional[str] = None,
    db: Optional[AsyncSession] = None,
) -> ExcelValidationResult:
    """
    Parse the uploaded 12-column Excel file and validate against business rules and DB.
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        return ExcelValidationResult(
            is_valid=False,
            total_rows=0,
            total_cartons=0,
            total_gross_weight=0.0,
            total_net_weight=0.0,
            total_quantity=0.0,
            line_summaries=[],
            rows=[],
            general_errors=[f"Failed to open Excel file: {str(e)}. Please upload a valid .xlsx file."],
        )

    sheet = wb.active
    if sheet is None:
        return ExcelValidationResult(
            is_valid=False,
            total_rows=0,
            total_cartons=0,
            total_gross_weight=0.0,
            total_net_weight=0.0,
            total_quantity=0.0,
            line_summaries=[],
            rows=[],
            general_errors=["The workbook contains no active sheets."],
        )

    # Find header row (first row with P/O or PO)
    header_row_idx = None
    col_map: Dict[str, int] = {}

    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if not row:
            continue
        row_str = " ".join([str(c or "").lower() for c in row])
        if "p/o" in row_str or "po #" in row_str or "po item" in row_str or "cart no" in row_str:
            header_row_idx = row_idx
            for c_idx, cell in enumerate(row):
                val = str(cell or "").strip().lower()
                if "p/o" in val or "po #" in val or val == "po" or val == "order number":
                    col_map["po_number"] = c_idx
                elif "po item" in val or "item" in val or "line" in val:
                    col_map["po_item"] = c_idx
                elif "pack no" in val or "slip" in val or "packing" in val:
                    col_map["pack_number"] = c_idx
                elif "cart no" in val or "box" in val or "carton" in val:
                    col_map["carton_number"] = c_idx
                elif "carton_ref" in val or "supplier_carton" in val or "roll" in val or "ref" in val:
                    col_map["supplier_carton_ref"] = c_idx
                elif "productcode" in val or "product code" in val or "material" in val or "item code" in val:
                    col_map["product_code"] = c_idx
                elif "lot" in val or "batch" in val:
                    col_map["lot_number"] = c_idx
                elif "width" in val or "dim" in val:
                    col_map["width"] = c_idx
                elif val == "gw" or "gross" in val:
                    col_map["gross_weight"] = c_idx
                elif val == "nw" or "net" in val:
                    col_map["net_weight"] = c_idx
                elif "quantity" in val or val == "qty":
                    col_map["quantity"] = c_idx
                elif "uom" in val or "unit" in val:
                    col_map["uom"] = c_idx
            break

    if header_row_idx is None:
        # Fallback to positional default (columns A to L)
        header_row_idx = 1
        col_map = {
            "po_number": 0, "po_item": 1, "pack_number": 2, "carton_number": 3,
            "supplier_carton_ref": 4, "product_code": 5, "lot_number": 6,
            "width": 7, "gross_weight": 8, "net_weight": 9, "quantity": 10, "uom": 11,
        }

    parsed_rows: List[ParsedCartonRow] = []
    unique_carton_ids = set()

    # Parse rows
    for row_idx, row in enumerate(sheet.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        def get_val(key: str) -> Any:
            idx = col_map.get(key)
            if idx is not None and idx < len(row):
                return row[idx]
            return None

        po_num = _clean_str(get_val("po_number"))
        po_item = _normalize_item_num(get_val("po_item"))

        # Skip rows without PO number
        if not po_num:
            continue

        carton_num = _parse_int(get_val("carton_number"))
        pack_num = _clean_str(get_val("pack_number"))
        carton_ref = _clean_str(get_val("supplier_carton_ref"))
        prod_code = _clean_str(get_val("product_code"))
        lot_num = _clean_str(get_val("lot_number"))
        width_val = _parse_float(get_val("width")) if get_val("width") is not None else None
        gw = _parse_float(get_val("gross_weight"))
        nw = _parse_float(get_val("net_weight"))
        qty = _parse_float(get_val("quantity"))
        uom = _clean_str(get_val("uom")) or "M"

        row_errors: List[str] = []

        if not po_item:
            row_errors.append("PO item is required")
        if qty <= 0:
            row_errors.append(f"Quantity must be greater than 0 (got {qty})")
        if gw < nw:
            row_errors.append(f"Gross weight ({gw}kg) cannot be less than Net weight ({nw}kg)")
        if nw <= 0:
            row_errors.append(f"Net weight must be > 0 (got {nw}kg)")
        if not prod_code:
            row_errors.append("ProductCode is required")

        parsed_row = ParsedCartonRow(
            row_index=row_idx,
            po_number=po_num,
            po_item=po_item,
            pack_number=pack_num,
            carton_number=carton_num,
            supplier_carton_ref=carton_ref,
            product_code=prod_code,
            lot_number=lot_num,
            width=width_val,
            gross_weight=gw,
            net_weight=nw,
            quantity=qty,
            uom=uom,
            errors=row_errors,
        )
        parsed_rows.append(parsed_row)

        carton_id = (po_num, pack_num, carton_num or carton_ref or row_idx)
        unique_carton_ids.add(carton_id)

    # Group by (po_number, po_item)
    groups: Dict[Tuple[str, str], List[ParsedCartonRow]] = {}
    for r in parsed_rows:
        key = (r.po_number, r.po_item)
        groups.setdefault(key, []).append(r)

    line_summaries: List[LineItemSummary] = []
    overall_valid = True

    # Database validation if session provided
    db_pos: Dict[str, PurchaseOrder] = {}
    db_available = False
    if db is not None:
        try:
            po_numbers = list({k[0] for k in groups.keys()})
            if po_numbers:
                stmt = select(PurchaseOrder).where(PurchaseOrder.po_number.in_(po_numbers))
                res = await db.execute(stmt)
                for po_obj in res.scalars().all():
                    db_pos[po_obj.po_number] = po_obj
                db_available = True
        except Exception:
            # In offline or dev testing, proceed without DB lookup
            db_available = False

    for (po_num, po_item), rows in groups.items():
        total_line_qty = sum(r.quantity for r in rows)
        total_gw = sum(r.gross_weight for r in rows)
        total_nw = sum(r.net_weight for r in rows)
        prod_code = rows[0].product_code
        group_errors: List[str] = []

        # Check row errors
        for r in rows:
            if r.errors:
                group_errors.extend([f"Row {r.row_index}: {err}" for err in r.errors])

        po_obj = db_pos.get(po_num)
        po_ordered_qty = None
        po_open_balance = None

        if db_available:
            if po_obj is None:
                group_errors.append(f"PO #{po_num} not found in database or belongs to another company")
            else:
                po_ordered_qty = float(po_obj.quantity or 0)
                po_open_balance = po_ordered_qty
                if total_line_qty > po_ordered_qty > 0:
                    group_errors.append(
                        f"Packed qty ({total_line_qty}) exceeds PO ordered quantity ({po_ordered_qty})"
                    )

        is_group_valid = len(group_errors) == 0
        if not is_group_valid:
            overall_valid = False

        summary = LineItemSummary(
            po_number=po_num,
            po_item=po_item,
            product_code=prod_code,
            total_qty=total_line_qty,
            carton_count=len(rows),
            total_gw=total_gw,
            total_nw=total_nw,
            po_ordered_qty=po_ordered_qty,
            po_open_balance=po_open_balance,
            is_valid=is_group_valid,
            errors=group_errors,
        )
        line_summaries.append(summary)

    total_gross = sum(r.gross_weight for r in parsed_rows)
    total_net = sum(r.net_weight for r in parsed_rows)
    total_qty = sum(r.quantity for r in parsed_rows)

    return ExcelValidationResult(
        is_valid=overall_valid and len(parsed_rows) > 0,
        total_rows=len(parsed_rows),
        total_cartons=len(unique_carton_ids),
        total_gross_weight=total_gross,
        total_net_weight=total_net,
        total_quantity=total_qty,
        line_summaries=line_summaries,
        rows=parsed_rows,
        general_errors=[] if parsed_rows else ["No valid carton rows were found in the uploaded file."],
    )


def generate_open_lines_template(open_lines: List[Dict[str, Any]]) -> bytes:
    """
    Generate a pre-filled Excel template (.xlsx) for selected open PO lines.
    Columns:
      A: P/O #
      B: PO item
      C: Pack No.
      D: Cart No
      E: Supplier_Carton_ref
      F: ProductCode
      G: Lot No.
      H: Width
      I: GW
      J: NW
      K: Quantity
      L: UOM
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Packing Sheet"

    headers = [
        "P/O #", "PO item", "Pack No.", "Cart No", "Supplier_Carton_ref",
        "ProductCode", "Lot No.", "Width", "GW", "NW", "Quantity", "UOM"
    ]
    ws.append(headers)

    # Style header row (Oniverse Navy / Emerald accent)
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for col_idx in range(1, 13):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    # Populate rows with open lines data
    for row_idx, item in enumerate(open_lines, start=2):
        row_data = [
            item.get("po_number", ""),
            item.get("po_item", "00100"),
            item.get("pack_number", ""),
            item.get("carton_number", ""),
            item.get("supplier_carton_ref", ""),
            item.get("product_code", ""),
            item.get("lot_number", ""),
            item.get("width", ""),
            item.get("gw", ""),
            item.get("nw", ""),
            item.get("quantity", ""),
            item.get("uom", "MTR"),
        ]
        ws.append(row_data)

        for col_idx in range(1, 13):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            # Pre-filled columns in soft tint
            if col_idx in (1, 2, 6, 12):
                cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            if col_idx in (4, 8, 9, 10, 11):
                cell.alignment = Alignment(horizontal="right")

    # Column widths
    col_widths = [15, 12, 16, 10, 24, 24, 14, 10, 10, 10, 12, 8]
    for idx, width in enumerate(col_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def generate_packing_template(open_lines: Optional[List[Dict[str, Any]]] = None) -> bytes:
    """Convenience alias for template generation with optional sample data."""
    if not open_lines:
        open_lines = [
            {
                "po_number": "2001297727",
                "po_item": "00100",
                "pack_number": "PACK-001",
                "carton_number": 1,
                "supplier_carton_ref": "SK104546-015.0-61851",
                "product_code": "ELST1K 000615",
                "lot_number": "LOT-2025-01",
                "width": 1.5,
                "gw": 25.5,
                "nw": 24.0,
                "quantity": 245,
                "uom": "M",
            },
            {
                "po_number": "2001297727",
                "po_item": "00100",
                "pack_number": "PACK-001",
                "carton_number": 2,
                "supplier_carton_ref": "SK104546-015.0-61852",
                "product_code": "ELST1K 000615",
                "lot_number": "LOT-2025-01",
                "width": 1.5,
                "gw": 26.0,
                "nw": 24.5,
                "quantity": 255,
                "uom": "M",
            },
        ]
    return generate_open_lines_template(open_lines)


# Exported alias
validate_packing_excel = parse_and_validate_packing_excel


